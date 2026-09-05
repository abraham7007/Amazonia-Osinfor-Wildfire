# -*- coding: utf-8 -*-
"""Versión editable en Word del cuaderno de campo. Mismo contenido que el PDF."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SALIDA = 'Cuaderno-de-campo-validacion-nodo.docx'
ANCHO  = 18.2          # cm útiles
CB     = '[    ]'
GRIS   = 'EAEAEA'
OSCURO = '404040'
LINEA  = '595959'

d = Document()
st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(8.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
st.paragraph_format.space_after = Pt(0); st.paragraph_format.space_before = Pt(0)
s = d.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin = Cm(1.9), Cm(1.3)
s.left_margin, s.right_margin = Cm(1.4), Cm(1.4)
s.header_distance = Cm(1.0)

# encabezado de página
h = s.header.paragraphs[0]
h.text = 'PROYECTO AMAZONÍA+  ·  PILOTO OSINFOR  ·  CCNN PAOYHAN (LORETO)\tCuaderno de campo — validación del nodo'
h.style.font.size = Pt(7.5); h.style.font.name = 'Arial'
for tab in (Cm(9),):
    pass
h.paragraph_format.tab_stops.add_tab_stop(Cm(18.2), WD_ALIGN_PARAGRAPH.RIGHT)
pf = s.footer.paragraphs[0]
pf.text = 'Rellenar a lápiz. No borrar: tachar con una línea, corregir al lado y poner iniciales.'
pf.style.font.size = Pt(7)


def _fijo(t, ancho_cm=ANCHO):
    """Word respeta los anchos solo con layout fijo."""
    t.autofit = False
    t.allow_autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), str(int(ancho_cm*567))); w.set(qn('w:type'),'dxa')
    tblPr.append(w)
    # la rejilla manda sobre el ancho de celda: hay que sincronizarla
    grid = t._tbl.find(qn('w:tblGrid'))
    if grid is not None:
        for gc, celda in zip(grid.findall(qn('w:gridCol')), t.rows[0].cells):
            if celda.width is not None:
                gc.set(qn('w:w'), str(int(celda.width.cm*567)))
    return t

def _bordes_celda(cell, lados):
    """lados: dict lado -> True (línea) / False (sin línea)"""
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for lado in ('top','left','bottom','right'):
        e = OxmlElement('w:'+lado)
        if lados.get(lado):
            e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),LINEA)
        else:
            e.set(qn('w:val'),'nil')
        b.append(e)
    tcPr.append(b)

def _sombra(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),color)
    tcPr.append(sh)

def _alto(row, cm, exacto=False):
    trPr = row._tr.get_or_add_trPr()
    hh = OxmlElement('w:trHeight'); hh.set(qn('w:val'), str(int(cm*567)))
    hh.set(qn('w:hRule'), 'exact' if exacto else 'atLeast'); trPr.append(hh)

def _txt(cell, texto, tam=8.5, negrita=False, color=None, centro=False):
    p = cell.paragraphs[0]
    if centro: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(texto); r.font.size = Pt(tam); r.font.name = 'Arial'; r.bold = negrita
    if color: r.font.color.rgb = color

def titulo(texto, sub=None):
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(texto); r.bold = True; r.font.size = Pt(13)
    if sub:
        q = d.add_paragraph(); q.paragraph_format.space_after = Pt(6)
        rr = q.add_run(sub); rr.font.size = Pt(7.5); rr.font.color.rgb = RGBColor(0x59,0x59,0x59)

def barra(texto):
    t = d.add_table(rows=1, cols=1); t.autofit = False
    c = t.rows[0].cells[0]; c.width = Cm(ANCHO)
    _sombra(c, OSCURO); _bordes_celda(c, {})
    _txt(c, texto, 9, True, RGBColor(0xFF,0xFF,0xFF))
    _alto(t.rows[0], 0.62)
    _fijo(t)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def campos(filas, anchos):
    """filas: listas de celdas; '' = espacio para escribir (con línea inferior)."""
    t = d.add_table(rows=len(filas), cols=len(anchos)); t.autofit = False
    for i, fila in enumerate(filas):
        _alto(t.rows[i], 0.72)
        for j, val in enumerate(fila):
            c = t.rows[i].cells[j]; c.width = Cm(anchos[j])
            _bordes_celda(c, {'bottom': val == ''})
            _txt(c, val)
    return _fijo(t)

def tabla(cabeceras, anchos, nfilas, alto=0.78):
    t = d.add_table(rows=nfilas+1, cols=len(cabeceras)); t.autofit = False
    _alto(t.rows[0], 0.95)
    for j, cab in enumerate(cabeceras):
        c = t.rows[0].cells[j]; c.width = Cm(anchos[j])
        _sombra(c, GRIS); _bordes_celda(c, {'top':1,'left':1,'bottom':1,'right':1})
        _txt(c, cab, 6.8, True, centro=True)
    for i in range(1, nfilas+1):
        _alto(t.rows[i], alto)
        for j in range(len(cabeceras)):
            c = t.rows[i].cells[j]; c.width = Cm(anchos[j])
            _bordes_celda(c, {'top':1,'left':1,'bottom':1,'right':1})
    return _fijo(t)

def renglones(n, alto=0.68):
    t = d.add_table(rows=n, cols=1); t.autofit = False
    for i in range(n):
        _alto(t.rows[i], alto)
        c = t.rows[i].cells[0]; c.width = Cm(ANCHO)
        _bordes_celda(c, {'bottom':1})
    return _fijo(t)

def cuadricula(alto_cm, paso=0.6):
    cols = int(ANCHO/paso); filas = int(alto_cm/paso)
    t = d.add_table(rows=filas, cols=cols); t.autofit = False
    for i in range(filas):
        _alto(t.rows[i], paso, exacto=True)
        for j in range(cols):
            c = t.rows[i].cells[j]; c.width = Cm(paso)
            _bordes_celda(c, {'top':1,'left':1,'bottom':1,'right':1})
    return _fijo(t)

def parrafo(texto, tam=8.5, esp=1.8):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(esp)
    r = p.add_run(texto); r.font.size = Pt(tam)
    return p

def aire(pt=4):
    d.add_paragraph().paragraph_format.space_after = Pt(pt)

def salto():
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ═══════════════ portada
titulo('Cuaderno de campo',
       'Validación del nodo de detección temprana de incendios (Edge AI + LoRaWAN)')
barra('DATOS DE LA CAMPAÑA')
campos([['Campaña N.º','','Del','','al','','Responsable','']],
       [2.4,1.6,1.0,2.6,0.8,2.6,2.4,4.8])
campos([['Integrantes','','Nodos (ID)','','Gateway (ID)','']],
       [2.2,5.8,2.4,3.4,2.6,1.8])
campos([['Versión de nodo  '+CB+' V1   '+CB+' V2   '+CB+' V3','','Firmware','','Modelo','']],
       [6.2,0.8,2.0,3.2,1.6,4.4])
aire()
barra('CÓMO LLENARLO')
for i in ['1.  Coordenadas en grados decimales, WGS84, 5 decimales (ej. −7.35812 / −74.90455). Hora local de 24 h.',
          '2.  Numerar las fotos de corrido durante toda la campaña y anotar el número. Sin ese número la foto no sirve de evidencia.',
          '3.  Si un dato no aplica, escribir n/a. A lápiz; para corregir, tachar con una línea y poner iniciales.',
          '4.  Los cuadros grises al pie de cada formato se llenan al terminar: ahí queda la conclusión.',
          '5.  Cada día se cierra firmando F-07.']:
    parrafo(i)
aire()
barra('CONTENIDO')
t = tabla(['','Formato','Qué cierra','Requisito','Copias'], [1.4,6.2,5.8,2.6,2.2], 7, alto=0.62)
datos = [['F-01','Punto de instalación y puesta en marcha','Dónde y cómo quedó cada nodo','003 · 006','1 por nodo'],
         ['F-02','Cobertura LoRaWAN (walk test)','Alcance real del enlace','003','2'],
         ['F-03','Ensayo de detección de humo','Tasa de detección y latencia','001 · 008','3'],
         ['F-04','Alertas y falsos positivos','Comportamiento continuo','001','4'],
         ['F-05','Energía y ambiente','Autonomía solar','004 · 005','4'],
         ['F-06','Inspección física e IP67','Sellado y montaje','006 · 007','2 por nodo'],
         ['F-07','Cierre de jornada','Trazabilidad diaria','—','1 por día']]
for i, fila in enumerate(datos, start=1):
    for j, v in enumerate(fila):
        _txt(t.rows[i].cells[j], v, 8, negrita=(j == 0))
aire(2)
parrafo('Los requisitos se citan abreviados: 001 = REQ-SYS-001, y así. '
        'Ver 01-requisitos/requisitos-de-sistema.md', 7.5)
salto()

# ═══════════════ F-01
titulo('F-01 · Punto de instalación y puesta en marcha', 'Una por nodo. Cierra REQ-SYS-003 y 006.')
barra('IDENTIFICACIÓN Y UBICACIÓN')
campos([['Nodo ID','','DevEUI','','Fecha','','Hora','']], [1.7,1.7,1.8,4.4,1.3,2.4,1.2,1.7])
campos([['Latitud','','Longitud','','Altitud (m)','','Precisión GPS (m)','']],
       [1.5,3.1,1.9,3.1,2.0,1.6,3.2,1.0])
campos([['Nombre del lugar / sector','','Cómo se llega','']], [4.2,4.0,2.6,7.4])
aire(2)
barra('MONTAJE Y CAMPO DE VISIÓN')
campos([['Soporte  '+CB+' árbol  '+CB+' poste  '+CB+' otro:','','Altura (m)','',
         'Azimut cámara (°)','','Inclinación (°)','']], [5.8,1.6,2.0,1.2,3.2,1.2,2.6,0.6])
campos([['Dosel  '+CB+' abierto  '+CB+' parcial  '+CB+' cerrado','',
         'Obstrucciones en el campo de visión','']], [6.6,0.6,5.4,5.6])
aire(2)
barra('ENLACE Y PUESTA EN MARCHA')
campos([['Distancia al gateway (m)','','Línea de vista  '+CB+' sí  '+CB+' no','',
         'Obstrucción principal','']], [4.2,1.6,4.2,0.6,3.8,3.8])
campos([['Hora de encendido','','Join OTAA  '+CB+' OK  '+CB+' falló','','SF','',
         'RSSI (dBm)','','SNR (dB)','']], [3.0,1.4,4.2,0.6,0.8,1.2,2.2,1.4,1.8,1.6])
campos([['Batería (V)','','SoC (%)','','Primer heartbeat, hora','',
         'Imagen de prueba  '+CB+' sí  '+CB+' no','']], [2.2,1.4,1.8,1.2,3.8,1.4,5.2,1.2])
aire(2)
barra('CROQUIS  (norte arriba · marcar nodo, campo de visión y gateway)')
cuadricula(8.4)
aire(2)
campos([['Fotos N.º  del','','al','','Observaciones','','Iniciales','']],
       [2.6,1.4,0.8,1.4,2.8,6.2,2.0,1.0])
salto()

# ═══════════════ F-02
titulo('F-02 · Cobertura LoRaWAN (walk test)',
       'Cierra REQ-SYS-003 (meta: 5 km). Alejarse del gateway tomando un punto cada 250–500 m. '
       'En cada punto enviar 10 tramas de prueba y anotar cuántas llegaron.')
campos([['Fecha','','Gateway ID','','Nodo usado','','Modo  '+CB+' ADR  '+CB+' SF fijo:','']],
       [1.3,2.0,2.2,2.0,2.4,2.0,4.4,1.9])
aire(2)
tabla(['N.º','Hora','Latitud','Longitud','Distancia al gateway (m)','SF','Llegaron de 10',
       'RSSI (dBm)','SNR (dB)','Observaciones'],
      [1.0,1.4,2.4,2.4,2.2,1.0,1.6,1.6,1.4,3.2], 15, alto=0.9)
aire(3)
barra('AL TERMINAR')
campos([['Distancia máxima con 9 o 10 de 10','','m','','¿Se llegó a 5 km?  '+CB+' sí  '+CB+' no','',
         'Iniciales','']], [5.2,1.6,0.6,0.6,5.4,1.4,2.0,1.4])
campos([['Qué limitó el alcance','']], [3.6,14.6])
renglones(1)
salto()

# ═══════════════ F-03
titulo('F-03 · Ensayo de detección de humo',
       'Cierra REQ-SYS-001 (meta: 85 % de detección) y REQ-SYS-008 (latencia). '
       'Generar humo controlado a distancias crecientes dentro del campo de visión.')
campos([['Fecha','','Nodo ID','','Fuente de humo (material)','','Umbral de confianza (%)','']],
       [1.3,2.0,1.8,1.8,4.4,3.0,4.2,1.2])
campos([['Cadencia de captura (min)','','Cielo','','Viento','','Temp. (°C)','']],
       [4.2,1.2,1.2,3.2,1.4,3.2,2.0,1.2])
aire(2)
tabla(['Ensayo N.º','Hora de inicio','Distancia al nodo (m)','Azimut (°)','¿Detectó? S / N',
       'Hora de la alerta','Confianza (%)','Foto N.º','Observaciones'],
      [1.4,1.8,2.2,1.4,1.8,2.0,2.0,1.2,4.2], 11, alto=1.06)
aire(3)
barra('AL TERMINAR')
campos([['Ensayos','','Detectados','','Tasa (%)','','Distancia máx. con detección (m)','',
         'Iniciales','']], [1.6,1.2,2.2,1.2,1.6,1.2,5.4,1.4,1.8,0.6])
campos([['¿Cumple el 85 %?  '+CB+' sí  '+CB+' no','','Notas sobre fallos','']],
       [4.2,1.0,3.2,9.8])
renglones(1)
salto()

# ═══════════════ F-04
titulo('F-04 · Alertas y falsos positivos',
       'Cierra REQ-SYS-001 en operación continua. Anotar toda alerta de humo, verdadera o falsa. '
       'Los falsos positivos por neblina y nube baja son el riesgo principal del modelo, y anotar '
       'su causa es lo que permite reentrenarlo.')
campos([['Nodo(s) en operación','','Del','','al','']], [4.0,5.2,1.0,3.2,0.8,4.0])
aire(2)
tabla(['Fecha','Hora','Nodo','Confianza (%)','¿Humo real? S / N / dudoso',
       'Si es falso, causa probable (neblina · nube · sol · polvo · insecto · otro)',
       'Foto N.º','Observaciones'],
      [1.8,1.4,1.4,1.8,2.6,5.4,1.2,2.6], 17, alto=0.94)
aire(3)
barra('AL TERMINAR')
campos([['Alertas de humo','','Verdaderas','','Falsas','','Falsos positivos (%)','',
         'Iniciales','']], [3.0,1.2,2.2,1.2,1.4,1.2,3.6,1.2,1.8,1.4])
campos([['Causa de falso positivo más frecuente','']], [5.8,12.4])
salto()

# ═══════════════ F-05
titulo('F-05 · Energía y ambiente',
       'Cierra REQ-SYS-004 (autonomía de 2–3 días sin sol) y REQ-SYS-005 (clima amazónico). '
       'Dos lecturas al día bastan: temprano en la mañana, antes de que cargue el panel, y al '
       'final de la tarde. La temperatura se toma del heartbeat.')
campos([['Nodo(s) / gateway','','Del','','al','']], [3.6,5.6,1.0,3.2,0.8,4.0])
aire(2)
tabla(['Fecha','Hora','Equipo nodo / GW','Batería (V)','SoC (%)','Temp. (°C)',
       'Cielo despejado / nubes / lluvia','¿Llovió anoche?','Observaciones'],
      [1.8,1.4,2.2,1.8,1.4,1.6,4.0,1.8,2.2], 17, alto=0.94)
aire(3)
barra('AL TERMINAR')
campos([['Días seguidos con poco sol','','¿Siguió operando?  '+CB+' sí  '+CB+' no','',
         'Tensión mínima (V)','','Iniciales','']], [4.4,1.2,4.4,1.0,3.2,1.2,1.8,1.0])
campos([['Notas (sombra sobre el panel, condensación, suciedad)','']], [7.4,10.8])
salto()

# ═══════════════ F-06 + F-07
titulo('F-06 · Inspección física e IP67',
       'Cierra REQ-SYS-006 y 007. Al instalar, tras la primera lluvia fuerte y al retirar.')
campos([['Fecha','','Nodo ID','','Momento  '+CB+' instalación  '+CB+' tras lluvia  '+CB+' retiro','']],
       [1.3,2.0,1.8,1.8,7.6,3.7])
aire(1)
items = ['Gabinete cerrado, tornillos completos y empaque en su sitio',
         'Sin condensación ni agua dentro',
         'Prensaestopas apretados y cables con curva de goteo',
         'Lente y visor limpios, sin gotas ni telarañas',
         'Fijación firme al soporte y antena vertical',
         'Panel solar limpio, sin sombra y bien orientado',
         'Sin fauna, insectos ni nidos en el gabinete',
         'Sin oxidación ni señales de manipulación']
t = tabla(['Punto de inspección','OK','Observado','Detalle / acción'], [8.4,1.2,1.8,6.8],
          len(items), alto=0.68)
for i, it in enumerate(items, start=1):
    _txt(t.rows[i].cells[0], it, 8.5)
    _txt(t.rows[i].cells[1], CB, 8.5, centro=True)
    _txt(t.rows[i].cells[2], CB, 8.5, centro=True)
aire(2)
campos([['Fotos N.º','','Iniciales','']], [2.0,4.0,2.0,10.2])
aire(3)

titulo('F-07 · Cierre de jornada', 'Una por día. Es lo que da trazabilidad a la campaña.')
campos([['Fecha','','Jornada N.º','','Clima del día','','Participantes','']],
       [1.3,2.0,2.4,1.2,2.6,3.2,2.4,3.1])
aire(1)
barra('QUÉ SE HIZO');                                 renglones(3, 0.58); aire(1)
barra('INCIDENCIAS Y DECISIONES TOMADAS EN CAMPO');   renglones(3, 0.58); aire(1)
barra('PENDIENTE PARA MAÑANA  ·  LECCIÓN APRENDIDA'); renglones(3, 0.58); aire(1)
cb = '[  ]'
campos([['Formatos llenados:  '+'   '.join(cb+' F-0'+str(i) for i in range(1,7)),' ']], [13.8,4.4])
campos([['Fotos del día, del','','al','','Firma del responsable','','Segunda firma','']],
       [3.0,1.4,0.9,1.3,4.0,3.4,2.6,1.6])

d.save(SALIDA)
print('generado:', SALIDA)
